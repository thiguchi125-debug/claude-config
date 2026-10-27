#!/usr/bin/env python3
"""後援会名簿（統合版xlsx）のあいさつ管理。Claude Code から呼ぶ。

  aisatsu.py status                         挨拶漏れを町別に表示（ファイルは変えない）
  aisatsu.py log <番号|氏名> <YYYY-MM-DD> <方法> <内容>   接触を1件記録して再計算
  aisatsu.py mark <◎|○|△|×> <番号,番号,...|町:中庄|G:B>   印を付けて再計算
  aisatsu.py refresh                        再計算だけ

書き込みは毎回 Drive 上の最新版を読み、保存直前にハッシュが変わっていないか確認してから同名で上書きする。
"""
import sys, os, hashlib, datetime as dt, tempfile, shutil
from collections import defaultdict
import openpyxl
from copy import copy
from openpyxl.styles import PatternFill, Font, Alignment

PATH = os.environ.get("AISATSU_PATH") or os.path.expanduser("~/Library/CloudStorage/GoogleDrive-t.higuchi125@gmail.com/マイドライブ/草川たくや 議会質問アーカイブ/日常資料アーカイブ/00_名簿・個人情報/後援会名簿_統合版_※個人情報注意_2026-10-27.xlsx")
INTERVAL = {"◎": 60, "○": 120, "△": 365}   # 何日あいたら「漏れ」か
SOON = 14                                   # 期限まで何日で「もうすぐ」か
RED, YEL, GRN, GRY = "F4CCCC", "FFF2CC", "D9EAD3", "EFEFEF"
# 名簿シートの列（1始まり）
C_NO, C_MARK, C_NAME, C_TOWN, C_LAST = 1, 2, 4, 6, 12
STATUS_COLS = ["次回の期限", "あいさつ状態"]

def norm_town(t):
    """「三寺町」と「三寺」を同じ町として数える"""
    t = (t or "").strip()
    if not t: return "（町不明）"
    return t[:-1] if t.endswith("町") and len(t) > 2 else t

def md5(p):
    return hashlib.md5(open(p, "rb").read()).hexdigest()

def today():
    return dt.date.today()

def load():
    h = md5(PATH)
    return openpyxl.load_workbook(PATH), h

def save(wb, h):
    if md5(PATH) != h:
        sys.exit("中止: 読み込んだ後に Drive 側でファイルが変わった。最新版から再実行すること。")
    fd, tmp = tempfile.mkstemp(suffix=".xlsx"); os.close(fd)
    wb.save(tmp); shutil.copyfile(tmp, PATH); os.remove(tmp)

def ensure_cols(ws):
    hdr = [c.value for c in ws[1]]
    cols = {}
    for name in STATUS_COLS:
        if name in hdr:
            cols[name] = hdr.index(name) + 1
        else:
            ws.cell(1, ws.max_column + 1, name)
            src = ws.cell(1, 1)
            new = ws.cell(1, ws.max_column)
            new.font = copy(src.font); new.fill = copy(src.fill)
            cols[name] = ws.max_column
    return cols

def as_date(v):
    if isinstance(v, dt.datetime): return v.date()
    if isinstance(v, dt.date): return v
    if isinstance(v, str) and v[:10].count("-") == 2:
        try: return dt.date.fromisoformat(v[:10])
        except ValueError: return None
    return None

def judge(mark, last):
    """→ (状態ラベル, 期限, 色, 超過日数)"""
    if mark not in INTERVAL:
        return ("対象外" if mark == "×" else "未判定", None, GRY, None)
    if last is None:
        return ("漏れ（未接触）", None, RED, 10**6)
    due = last + dt.timedelta(days=INTERVAL[mark])
    left = (due - today()).days
    if left < 0: return (f"漏れ（{-left}日超過）", due, RED, -left)
    if left <= SOON: return (f"もうすぐ（あと{left}日）", due, YEL, None)
    return ("OK", due, GRN, None)

def recalc(wb):
    ws = wb["名簿"]
    # 接触履歴から最終接触日を取り直す
    last = {}
    if "接触履歴" in wb.sheetnames:
        for r in wb["接触履歴"].iter_rows(min_row=2, values_only=True):
            d = as_date(r[0]); no = r[1]
            if d and no and (no not in last or d > last[no]): last[no] = d
    cols = ensure_cols(ws)
    people = []
    for row in ws.iter_rows(min_row=2):
        no = row[C_NO-1].value
        if not no: continue
        d = max(filter(None, [last.get(no), as_date(row[C_LAST-1].value)]), default=None)
        row[C_LAST-1].value = d.isoformat() if d else None
        mark = (row[C_MARK-1].value or "").strip()
        label, due, color, over = judge(mark, d)
        ws.cell(row[0].row, cols["次回の期限"], due.isoformat() if due else None)
        ws.cell(row[0].row, cols["あいさつ状態"], label)
        f = PatternFill("solid", fgColor=color)
        for c in (row[C_MARK-1], row[C_NAME-1], ws.cell(row[0].row, cols["あいさつ状態"])): c.fill = f
        people.append(dict(no=no, name=row[C_NAME-1].value, town=norm_town(row[C_TOWN-1].value),
                           mark=mark, label=label, color=color, over=over, last=d))
    build_board(wb, people)
    return people

def build_board(wb, people):
    if "挨拶漏れ" in wb.sheetnames: del wb["挨拶漏れ"]
    ws = wb.create_sheet("挨拶漏れ", 0)
    wb.active = 0
    tgt = [p for p in people if p["mark"] in INTERVAL]
    red = [p for p in tgt if p["color"] == RED]; yel = [p for p in tgt if p["color"] == YEL]
    ws.append([f"挨拶漏れボード（{today().isoformat()} 時点・Claude Code が更新）"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"対象 {len(tgt)}名（◎{INTERVAL['◎']}日／○{INTERVAL['○']}日／△{INTERVAL['△']}日あいたら漏れ）　"
               f"漏れ {len(red)}名　もうすぐ {len(yel)}名　OK {len(tgt)-len(red)-len(yel)}名　未判定 {sum(1 for p in people if not p['mark'])}名"])
    ws.append([])
    ws.append(["町・エリア", "漏れ", "もうすぐ", "OK", "漏れている人（番号 氏名・印）"])
    for c in ws[4]: c.font = Font(bold=True)
    by = defaultdict(list)
    for p in tgt: by[p["town"]].append(p)
    def key(item):
        ps = item[1]; return (-sum(p["color"] == RED for p in ps), -sum(p["color"] == YEL for p in ps), item[0])
    for town, ps in sorted(by.items(), key=key):
        r = [p for p in ps if p["color"] == RED]; y = [p for p in ps if p["color"] == YEL]
        ws.append([town, len(r), len(y), len(ps)-len(r)-len(y),
                   "、".join(f"{p['no']} {p['name']}{p['mark']}" for p in sorted(r, key=lambda p: p["no"]))])
        color = RED if r else (YEL if y else GRN)
        for c in ws[ws.max_row]: c.fill = PatternFill("solid", fgColor=color)
        ws.cell(ws.max_row, 5).alignment = Alignment(wrap_text=True, vertical="top")
    for col, w in zip("ABCDE", [16, 7, 9, 6, 90]): ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"

def find(ws, key):
    hits = [r for r in ws.iter_rows(min_row=2) if r[C_NO-1].value == key or r[C_NAME-1].value == key]
    if not hits:
        hits = [r for r in ws.iter_rows(min_row=2) if r[C_NAME-1].value and key in r[C_NAME-1].value]
    return hits

def report(people):
    tgt = [p for p in people if p["mark"] in INTERVAL]
    red = [p for p in tgt if p["color"] == RED]; yel = [p for p in tgt if p["color"] == YEL]
    print(f"対象{len(tgt)}名／漏れ{len(red)}／もうすぐ{len(yel)}／未判定{sum(1 for p in people if not p['mark'])}")
    by = defaultdict(list)
    for p in red + yel: by[p["town"]].append(p)
    for town, ps in sorted(by.items(), key=lambda x: -len(x[1])):
        print(f"■{town}（{len(ps)}）: " + "、".join(f"{p['no']} {p['name']}{p['mark']}[{p['label']}]" for p in ps))

def main():
    a = sys.argv[1:]
    if not a or a[0] not in ("status", "log", "mark", "refresh"): sys.exit(__doc__)
    wb, h = load()
    if a[0] == "status":
        report(recalc(wb)); return
    if a[0] == "log":
        if len(a) < 5: sys.exit(__doc__)
        key, day, how, what = a[1], a[2], a[3], " ".join(a[4:])
        try: dt.date.fromisoformat(day)
        except ValueError: sys.exit(f"中止: 日付は YYYY-MM-DD で（受け取った値: {day}）")
        hits = find(wb["名簿"], key)
        if len(hits) != 1:
            sys.exit(f"中止: 「{key}」に該当 {len(hits)}件 → " + "、".join(f"{r[0].value} {r[C_NAME-1].value}" for r in hits))
        no, name = hits[0][C_NO-1].value, hits[0][C_NAME-1].value
        wb["接触履歴"].append([day, no, name, how, what])
        print(f"記録: {day} {no} {name} {how} {what}")
    if a[0] == "mark":
        mark, targets = a[1], " ".join(a[2:])
        if mark not in ("◎", "○", "△", "×", "-"): sys.exit("印は ◎ ○ △ × -（空欄に戻す）")
        ws = wb["名簿"]; n = 0
        for row in ws.iter_rows(min_row=2):
            no = row[C_NO-1].value
            if not no: continue
            t = targets
            ok = (t.startswith("町:") and t[2:] in (row[C_TOWN-1].value or "")) or \
                 (t.startswith("G:") and no.startswith(t[2:])) or \
                 (no in [x.strip() for x in t.replace("、", ",").split(",")])
            if ok: row[C_MARK-1].value = None if mark == "-" else mark; n += 1
        print(f"{n}名を {mark} に")
    people = recalc(wb)
    save(wb, h)
    report(people)

if __name__ == "__main__":
    main()
